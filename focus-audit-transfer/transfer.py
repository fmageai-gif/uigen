#!/usr/bin/env python3
"""Copy daily audit rows from ADHOC.xlsx into the Focus Audit SharePoint list.

Typical use on Windows:

    inspect.bat          once, to capture the form's fields and dropdown options
    run.bat --dry-run    fills one entry and stops before Save so you can check it
    run.bat              does the real thing

The browser profile lives in .browser-profile/ next to this script, so you sign
in to SharePoint once and stay signed in on later runs.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from playwright.sync_api import Page, TimeoutError as PWTimeout, sync_playwright

VERSION = "1.7"

HERE = Path(__file__).resolve().parent
CONFIG_PATH = HERE / "config.json"
PROFILE_DIR = HERE / ".browser-profile"
SUBMITTED_LOG = HERE / "submitted.json"
SCHEMA_OUT = HERE / "form_schema.json"
RUN_LOG = HERE / "last_run.log"
SHOTS = HERE / "screenshots"

# Form field label -> how we fill it. Labels must match the SharePoint form exactly.
TEXT_FIELDS = ("CaseID", "Call/Chat ID", "Case Subject", "What are the Opportunities", "Comments/Summary")
CHOICE_FIELDS = ("LOB", "Call/Chat Selection Criteria", "Suggested Resolution Code", "Validation")
DATE_FIELDS = ("Call Listening Date", "Call/Chat Date")
PERSON_FIELD = "Agent Name"

# Only these two markers are validity tags. Parentheses in expectedResolutionCode
# (e.g. "Subscription Cancellation (SubCan)") are part of the real code name and
# must never be stripped.
VALIDITY_SUFFIX = re.compile(r"\s*\((?:Pass|Invalid)\)\s*$", re.IGNORECASE)


# --------------------------------------------------------------------------- #
# Config + logging
# --------------------------------------------------------------------------- #

def load_config() -> dict[str, Any]:
    if not CONFIG_PATH.exists():
        sys.exit(f"Missing {CONFIG_PATH.name} - it must sit next to transfer.py")
    return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))


_log_lines: list[str] = []


def log(msg: str) -> None:
    print(msg, flush=True)
    _log_lines.append(f"{datetime.now():%H:%M:%S}  {msg}")


def flush_log() -> None:
    RUN_LOG.write_text("\n".join(_log_lines), encoding="utf-8")


def snap(page: Page, name: str) -> None:
    """Save a screenshot plus the page HTML, so a failure can be diagnosed later."""
    try:
        SHOTS.mkdir(exist_ok=True)
        stamp = f"{datetime.now():%Y%m%d-%H%M%S}-{name}"
        page.screenshot(path=str(SHOTS / f"{stamp}.png"), full_page=True)
        (SHOTS / f"{stamp}.html").write_text(page.content(), encoding="utf-8")
        log(f"    saved screenshots/{stamp}.png")
    except Exception as exc:
        log(f"    (could not capture screenshot: {exc})")


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #

@dataclass
class Row:
    excel_row: int
    audit_id: str = ""
    audit_date: str = ""
    call_date: str = ""
    agent_name: str = ""
    agent_email: str = ""
    lob: str = ""
    case_number: str = ""
    case_subject: str = ""
    genesys_id: str = ""
    validation: str = ""
    case_resolution: str = ""
    expected_resolution: str = ""
    remarks: str = ""
    auditor_name: str = ""
    quick_case: str = ""

    @property
    def is_invalid(self) -> bool:
        return self.validation.strip().lower() == "invalid"

    @property
    def suggested_resolution_code(self) -> str:
        """Invalid audits take the tagged code (col R); valid ones take the expected code (col S).

        The trailing "(Invalid)" / "(Pass)" marker is dropped, because it labels the
        audit outcome rather than forming part of the resolution code itself.
        """
        source = self.case_resolution if self.is_invalid else self.expected_resolution
        return VALIDITY_SUFFIX.sub("", source).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.month}/{value.day}/{value.year}"
    if isinstance(value, date):
        return f"{value.month}/{value.day}/{value.year}"
    return str(value).strip()


def _as_date(text: str) -> date | None:
    text = text.strip()
    if not text:
        return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def read_rows(cfg: dict[str, Any], xlsx: Path) -> list[Row]:
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    sheet = cfg["sheet_name"]
    if sheet not in wb.sheetnames:
        sys.exit(f"Sheet '{sheet}' not found. Sheets present: {', '.join(wb.sheetnames)}")
    ws = wb[sheet]

    idx = {name: column_index_from_string(letter) - 1 for name, letter in cfg["columns"].items()}
    widest = max(idx.values()) + 1

    rows: list[Row] = []
    for n, values in enumerate(ws.iter_rows(min_row=cfg["header_row"] + 1, values_only=True),
                               start=cfg["header_row"] + 1):
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        padded = list(values) + [None] * (widest - len(values))
        row = Row(excel_row=n, **{name: _cell_text(padded[i]) for name, i in idx.items()})
        if row.audit_id or row.case_number:
            rows.append(row)
    wb.close()
    return rows


def select_rows(cfg: dict[str, Any], rows: list[Row], args: argparse.Namespace) -> list[Row]:
    """Narrow to this auditor's rows, within the requested dates, minus anything already sent."""
    auditor = cfg["auditor_name"].strip().lower()
    picked = [r for r in rows if r.auditor_name.strip().lower() == auditor]
    log(f"{len(picked)} of {len(rows)} rows belong to {cfg['auditor_name']}")

    if args.date:
        target = _as_date(args.date) or date.fromisoformat(args.date)
        picked = [r for r in picked if _as_date(r.audit_date) == target]
        log(f"{len(picked)} with auditDate {target:%m/%d/%Y}")
    elif not args.all:
        cutoff = date.today().toordinal() - args.last_days
        picked = [r for r in picked
                  if (d := _as_date(r.audit_date)) and d.toordinal() >= cutoff]
        log(f"{len(picked)} within the last {args.last_days} days")

    if not args.force:
        done = set(load_submitted())
        before = len(picked)
        picked = [r for r in picked if r.audit_id not in done]
        if before != len(picked):
            log(f"{before - len(picked)} already submitted on an earlier run - skipping")

    return picked[: args.limit] if args.limit else picked


def _row_choice_value(row: Row, label: str) -> str:
    return {
        "Call/Chat Selection Criteria": row.expected_resolution,
        "Suggested Resolution Code": row.suggested_resolution_code,
        "LOB": row.lob,
        "Validation": row.validation,
    }.get(label, "")


def partition_rows(cfg: dict[str, Any], rows: list[Row]) -> tuple[list[Row], list[tuple[Row, str, str]]]:
    """Split off rows that cannot be filled in, rather than failing on them mid-run.

    Both Call/Chat Selection Criteria and Suggested Resolution Code are required
    and offer a fixed set of options that the workbook's resolution codes do not
    fully cover. A row with no match can never be saved, so it is identified up
    front and reported with the field and value that could not be placed.
    """
    required = cfg.get("required_choice_options") or {}
    if not required:
        return rows, []

    keep: list[Row] = []
    skip: list[tuple[Row, str, str]] = []
    for row in rows:
        for label, options in required.items():
            if not options:
                continue
            raw = _row_choice_value(row, label)
            value = cfg.get("choice_map", {}).get(label, {}).get(raw, raw).strip()
            if not _best_option(value, options):
                skip.append((row, label, raw))
                break
        else:
            keep.append(row)
    return keep, skip


def describe_row(cfg: dict[str, Any], row: Row) -> dict[str, str]:
    """What each dropdown would actually receive for this row, after mapping."""
    required = cfg.get("required_choice_options") or {}
    out: dict[str, str] = {}
    for label in CHOICE_FIELDS:
        raw = _row_choice_value(row, label)
        value = cfg.get("choice_map", {}).get(label, {}).get(raw, raw).strip()
        options = required.get(label) or []
        landed = _best_option(value, options) if options else value
        out[label] = landed or "<<no match>>"
    return out


def run_preview(cfg: dict[str, Any], rows: list[Row],
                skipped: list[tuple[Row, str, str]]) -> None:
    """Show what would be submitted, without opening a browser or touching the list."""
    log(f"\n{'=' * 78}\nWOULD SUBMIT ({len(rows)} rows)\n{'=' * 78}")
    for row in rows:
        landed = describe_row(cfg, row)
        log(f"\n{row.audit_id}   audit {row.audit_date}   call {row.call_date}   "
            f"case {row.case_number}")
        log(f"    Agent Name                    {row.agent_email}")
        for label, value in landed.items():
            source = _row_choice_value(row, label)
            arrow = "" if value == source else f"      <- {source!r}"
            log(f"    {label:29} {value}{arrow}")
        log(f"    Call/Chat ID                  {row.genesys_id}")
        log(f"    Case Subject                  {row.case_subject[:52]}")
        log(f"    Comments/Summary              {row.remarks[:52]}")

    log(f"\n{'=' * 78}\nWOULD SKIP ({len(skipped)} rows)\n{'=' * 78}")
    for row, label, value in skipped:
        log(f"  {row.audit_id}  case {row.case_number:12}  {label} <- {value!r}")

    total = len(rows) + len(skipped)
    if total:
        log(f"\n{len(rows)} of {total} rows would transfer "
            f"({100 * len(rows) // total}%).")
    log("Nothing was submitted - this was a preview.")


def load_submitted() -> list[str]:
    if not SUBMITTED_LOG.exists():
        return []
    try:
        return json.loads(SUBMITTED_LOG.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return []


def mark_submitted(audit_id: str) -> None:
    done = load_submitted()
    if audit_id and audit_id not in done:
        done.append(audit_id)
        SUBMITTED_LOG.write_text(json.dumps(done, indent=2), encoding="utf-8")


# --------------------------------------------------------------------------- #
# SharePoint form
# --------------------------------------------------------------------------- #

class FormError(RuntimeError):
    pass


@dataclass
class Form:
    """Thin wrapper over the modern SharePoint 'New item' panel.

    Locators are built around one observation from the live form: every field's
    accessible name *starts with* the field title and then adds boilerplate, e.g.

        "CaseID Required Field, empty, field editor. "
        "Call Listening Date, Tab to navigate through the date time callout, ..."

    So a starts-with match on aria-label pins the right control whatever its type.
    The HTML ids are no help - both date fields share one id, and the text fields
    use React counters (TextField19, TextField22, ...) that shift as the form changes.
    """

    page: Page
    choice_map: dict[str, dict[str, str]] = field(default_factory=dict)

    # -- locating ---------------------------------------------------------- #

    def _control(self, label: str, prefer: str = "field"):
        """Find a field's control by the prefix of its accessible name.

        Each editor is wrapped in a span repeating the control's own aria-label, so
        matching on aria-label alone lands on the wrapper - which cannot be typed
        into. Real inputs are therefore matched first, then comboboxes (the choice
        fields are divs), and only then anything else carrying the name.
        """
        name = label.replace('"', '\\"')
        editable = (f'input[aria-label^="{name}"]:visible, '
                    f'textarea[aria-label^="{name}"]:visible, '
                    f'[aria-label^="{name}"] input:visible, '
                    f'[aria-label^="{name}"] textarea:visible')
        combobox = f'[role="combobox"][aria-label^="{name}"]:visible'
        anything = f'[aria-label^="{name}"]:visible'

        order = (combobox, editable, anything) if prefer == "combobox" else (editable, combobox, anything)
        for selector in order:
            candidate = self.page.locator(selector)
            if candidate.count():
                return candidate.first

        candidate = self.page.get_by_label(label, exact=False)
        if candidate.count():
            return candidate.first

        quoted = _xpath_literal(label)
        control = "self::input or self::textarea or @role='combobox'"
        candidate = self.page.locator(
            f"xpath=//*[normalize-space(text())={quoted}]/following::*[{control}][1]")
        if candidate.count():
            return candidate.first

        raise FormError(f"Could not find the '{label}' field on the form")

    # -- filling ----------------------------------------------------------- #

    def set_text(self, label: str, value: str) -> None:
        control = self._control(label)
        control.click()
        control.fill(value)

    def set_date(self, label: str, value: str, time_of_day: str) -> None:
        """Fill a SharePoint date+time field and confirm the time came out right.

        The control is a single combined text field that displays "8/1/2026 12:00 AM",
        so the date and time are typed together and the callout is closed with Enter.
        The time defaults to 12:00 AM, but it is verified rather than assumed.
        """
        control = self._control(label)
        control.click()
        control.fill(f"{value} {time_of_day}")
        control.press("Enter")
        self.page.wait_for_timeout(400)

        landed = (control.input_value() or "").strip()
        if _same_time(landed, time_of_day) and value.lstrip("0") in landed.replace(" 0", " "):
            log(f"    {label}: {landed}")
            return

        # Typing them together was rejected - fall back to date only, then the
        # time control inside the open callout.
        control.click()
        control.fill(value)
        self._set_time_in_callout(time_of_day)
        control.press("Enter")
        self.page.wait_for_timeout(400)

        landed = (control.input_value() or "").strip()
        log(f"    {label}: {landed}")
        if not _same_time(landed, time_of_day):
            raise FormError(
                f"{label} ended up as '{landed}' - expected the time to be {time_of_day}")

    def _set_time_in_callout(self, time_of_day: str) -> None:
        """Set the time inside the open date callout, without touching the date box."""
        candidates = self.page.locator(
            "[class*='TimePicker'] input:visible, [class*='timePicker'] input:visible, "
            "[class*='TimePicker'] [role='combobox']:visible")
        if not candidates.count():
            return
        control = candidates.first
        try:
            control.click()
            option = self.page.get_by_role("option", name=time_of_day, exact=False)
            if option.count():
                option.first.click()
            else:
                control.fill(time_of_day)
        except PWTimeout:
            log(f"    ! could not set the time to {time_of_day}")

    def set_person(self, label: str, email: str) -> None:
        control = self._control(label)
        control.click()
        control.fill(email)

        # Only visible suggestions count: the form keeps other, closed listboxes in
        # the DOM. "Search Directory" is an action at the foot of the list, not a person.
        suggestion = self.page.locator(
            ".ms-Suggestions-item:visible, [class*='peoplePicker'] [role='option']:visible, "
            "[role='listbox'] [role='option']:visible, [class*='suggestionItem']:visible")
        try:
            suggestion.first.wait_for(state="visible", timeout=15_000)
        except PWTimeout:
            raise FormError(f"No directory match appeared for {email}")

        texts = suggestion.all_inner_texts()
        pick = 0
        for i, text in enumerate(texts):
            if "search directory" in text.strip().lower():
                continue
            pick = i
            break
        else:
            raise FormError(f"Only 'Search Directory' was offered for {email}")

        suggestion.nth(pick).click()
        self.page.wait_for_timeout(300)
        if not self._person_pill_present():
            raise FormError(f"Picked a suggestion for {email} but no name pill appeared")
        log(f"    {label}: {texts[pick].splitlines()[0].strip()}")

    def _person_pill_present(self) -> bool:
        pill = self.page.locator(
            "[class*='personaPill']:visible, [class*='pickerItem']:visible, "
            "[class*='ms-PickerPersona']:visible, [class*='personDisplayPill']:visible")
        return bool(pill.count())

    def close_dropdown(self) -> None:
        """Make sure an open dropdown has really gone away before touching the next field.

        These dropdowns render into a portal that covers the fields below them, so
        a leftover open list silently swallows the next click - the failure shows up
        as an unrelated field timing out with "subtree intercepts pointer events".
        """
        for _ in range(3):
            if not self.page.get_by_role("option").count():
                return
            self.page.keyboard.press("Escape")
            self.page.wait_for_timeout(400)

        # Escape did not take: click a neutral part of the panel instead.
        try:
            self.page.get_by_text("New item", exact=False).first.click(timeout=3_000)
            self.page.wait_for_timeout(300)
        except Exception:
            pass
        if self.page.get_by_role("option").count():
            log("    ! a dropdown is still open; the next field may be blocked")

    def set_choice(self, label: str, value: str, interactive: bool) -> None:
        """Pick an option from one of the filter-style choice dropdowns."""
        wanted = self.choice_map.get(label, {}).get(value, value).strip()
        control = self._control(label, prefer="combobox")
        control.click()

        options = self.page.get_by_role("option")
        try:
            options.first.wait_for(state="visible", timeout=10_000)
        except PWTimeout:
            self.close_dropdown()
            raise FormError(f"The '{label}' dropdown did not open")

        texts = [t.strip() for t in options.all_inner_texts()]
        match = _best_option(wanted, texts)
        if match is None:
            # Escape must go to the page: clicking the field swaps the element out,
            # so pressing a key on the original locator would fail.
            self.close_dropdown()
            self._resolve_manually(label, value, wanted, texts, interactive)
            return

        options.nth(texts.index(match)).click()
        self.page.wait_for_timeout(200)
        self.close_dropdown()
        if match != wanted:
            log(f"    {label}: '{value}' -> '{match}'")
        else:
            log(f"    {label}: {match}")

    def _resolve_manually(self, label: str, excel_value: str, wanted: str,
                          options: list[str], interactive: bool) -> None:
        log(f"    ! No '{label}' option matches Excel value '{excel_value}'")
        log(f"      Available: {', '.join(options) or '(none found)'}")
        if not interactive:
            raise FormError(f"Unmatched {label} value '{excel_value}'")
        log(f"      Pick the right option in the browser yourself, then press Enter here.")
        input("      Press Enter once you've set it... ")
        log(f"      Tip: add \"{wanted}\": \"<the option you picked>\" under "
            f"choice_map -> {label} in config.json so it's automatic next time.")

    # -- submitting -------------------------------------------------------- #

    def save(self) -> None:
        self.page.get_by_role("button", name="Save", exact=True).click()
        self.page.wait_for_timeout(2500)


def _same_time(landed: str, time_of_day: str) -> bool:
    squash = lambda s: re.sub(r"\s+", "", s).lower()
    return squash(time_of_day) in squash(landed)


def _xpath_literal(text: str) -> str:
    """Quote a string for XPath 1.0, which has no escape syntax of its own."""
    if "'" not in text:
        return f"'{text}'"
    if '"' not in text:
        return f'"{text}"'
    parts = "', \"'\", '".join(text.split("'"))
    return f"concat('{parts}')"


def _best_option(wanted: str, options: list[str]) -> str | None:
    """Match an Excel value to a dropdown option.

    The two sides are worded differently in both directions: Excel's
    "Remote Solution (HW/SW Resolved Remotely via Phone/Chat)" has to reach the
    option "Remote Solution", while a stripped "Subscription Cancellation" has to
    reach "Subscription Cancellation (SubCan)". Exact matches always win first, so
    "Instant Ink" can never be dragged onto "Instant Ink Chat".
    """
    if not wanted:
        return None
    if wanted in options:
        return wanted

    lowered = wanted.lower()
    for opt in options:
        if opt.lower() == lowered:
            return opt

    # Option is the shorter, more general form of the Excel value.
    contained = [o for o in options if o and lowered.startswith(o.lower())]
    if contained:
        return max(contained, key=len)

    # Excel value is the shorter form; only accept it when it is unambiguous.
    expands = [o for o in options if o.lower().startswith(lowered)]
    return expands[0] if len(expands) == 1 else None


def open_new_item(page: Page, cfg: dict[str, Any]) -> Form:
    """Open the New item panel, trying the command bar then the direct form URL."""
    page.goto(cfg["list_url"], wait_until="domcontentloaded")
    page.wait_for_timeout(2500)

    attempts = (
        lambda: page.get_by_role("menuitem", name=cfg["new_item_button"], exact=False).first.click(timeout=8_000),
        lambda: page.locator("[data-automationid='newCommand']").first.click(timeout=8_000),
        lambda: page.get_by_role("button", name=cfg["new_item_button"], exact=False).first.click(timeout=8_000),
        lambda: page.locator("button:has-text('New'), a:has-text('New')").first.click(timeout=8_000),
    )
    for attempt in attempts:
        try:
            attempt()
            page.get_by_text("New item", exact=False).first.wait_for(state="visible", timeout=10_000)
            return Form(page=page, choice_map=cfg.get("choice_map", {}))
        except (PWTimeout, Exception):
            continue

    # Last resort: the list's own new-item form URL.
    new_form = cfg.get("new_form_url", "").strip()
    if new_form:
        page.goto(new_form, wait_until="domcontentloaded")
        try:
            page.get_by_text("New item", exact=False).first.wait_for(state="visible", timeout=20_000)
            return Form(page=page, choice_map=cfg.get("choice_map", {}))
        except PWTimeout:
            pass

    raise FormError("Could not open the New item panel")


def fill_row(form: Form, row: Row, cfg: dict[str, Any], interactive: bool) -> None:
    time_of_day = cfg.get("time_of_day", "12:00 AM")
    form.set_choice("LOB", row.lob, interactive)
    form.set_date("Call Listening Date", row.audit_date, time_of_day)
    form.set_person(PERSON_FIELD, row.agent_email)
    form.set_choice("Call/Chat Selection Criteria", row.expected_resolution, interactive)
    form.set_choice("Suggested Resolution Code", row.suggested_resolution_code, interactive)
    form.set_date("Call/Chat Date", row.call_date, time_of_day)
    form.set_text("CaseID", row.case_number)
    form.set_text("Call/Chat ID", row.genesys_id)
    form.set_text("Case Subject", row.case_subject)
    form.set_text("What are the Opportunities", cfg.get("opportunities_text", "N/A"))
    form.set_text("Comments/Summary", row.remarks)
    form.set_choice("Validation", row.validation, interactive)


# --------------------------------------------------------------------------- #
# Inspect mode
# --------------------------------------------------------------------------- #

def inspect_form(page: Page, cfg: dict[str, Any]) -> None:
    """Dump every field and every dropdown option so the mapping can be finalised.

    This is the step that unblocks everything else, so it never gives up: if the
    form cannot even be opened it still writes whatever it saw, plus screenshots.
    """
    schema: dict[str, Any] = {"tool_version": VERSION,
                              "captured_at": datetime.now().isoformat(), "fields": []}
    try:
        open_new_item(page, cfg)
    except Exception as exc:
        schema["open_error"] = str(exc)
        log(f"! Could not open the New item form: {exc}")
        log("  Open it by hand in the browser window - click New on the list.")
        input("  Press Enter once the New item panel is showing... ")

    # span[aria-label] matters: the form wraps each editor in one repeating the
    # control's own name, and a locator matching aria-label alone hits the wrapper.
    controls = page.locator("input, textarea, [role='combobox'], "
                            "[role='button'][aria-haspopup], span[aria-label]")
    for i in range(controls.count()):
        el = controls.nth(i)
        try:
            schema["fields"].append({
                "tag": el.evaluate("e => e.tagName.toLowerCase()"),
                "role": el.get_attribute("role"),
                "type": el.get_attribute("type"),
                "aria_label": el.get_attribute("aria-label"),
                "placeholder": el.get_attribute("placeholder"),
                "id": el.get_attribute("id"),
                "class": (el.get_attribute("class") or "")[:120],
            })
        except Exception as exc:  # a control can detach while the panel settles
            schema["fields"].append({"error": str(exc)})

    schema["choices"] = {}
    for label in CHOICE_FIELDS:
        captured: Any = {"error": "not reached"}
        try:
            form = Form(page=page)
            control = form._control(label, prefer="combobox")
            control.click()
            page.get_by_role("option").first.wait_for(state="visible", timeout=10_000)
            # Store before closing: tidying up must never lose what was captured.
            captured = [t.strip() for t in page.get_by_role("option").all_inner_texts()]
            log(f"  {label}: {len(captured)} options")
        except Exception as exc:
            captured = {"error": str(exc)}
            log(f"  ! {label}: {str(exc).splitlines()[0]}")
        finally:
            schema["choices"][label] = captured
            try:
                Form(page=page).close_dropdown()
            except Exception:
                pass

    schema["date_callout"] = {}
    for label in DATE_FIELDS:
        try:
            form = Form(page=page)
            control = form._control(label)
            control.click()
            page.wait_for_timeout(800)
            inner = page.locator("[class*='Callout']:visible input, [class*='callout']:visible input, "
                                 "[class*='TimePicker']:visible, [class*='timePicker']:visible")
            schema["date_callout"][label] = [{
                "tag": inner.nth(j).evaluate("e => e.tagName.toLowerCase()"),
                "aria_label": inner.nth(j).get_attribute("aria-label"),
                "id": inner.nth(j).get_attribute("id"),
                "value": inner.nth(j).get_attribute("value"),
                "class": (inner.nth(j).get_attribute("class") or "")[:120],
            } for j in range(min(inner.count(), 12))]
            log(f"  {label} callout: {len(schema['date_callout'][label])} controls")
        except Exception as exc:
            schema["date_callout"][label] = {"error": str(exc)}
        finally:
            try:
                page.keyboard.press("Escape")
                page.wait_for_timeout(300)
            except Exception:
                pass

    schema["labels_on_page"] = sorted({
        t.strip() for t in page.locator("label, [role='heading'], .ms-Label").all_inner_texts()
        if t.strip()
    })

    SCHEMA_OUT.write_text(json.dumps(schema, indent=2), encoding="utf-8")
    snap(page, "inspect-final")
    log(f"\nWrote {SCHEMA_OUT.name} - send this file back to finish the mapping.")
    log(f"Also send the newest screenshots/*.png if anything looked wrong.")

# --------------------------------------------------------------------------- #
# Excel retrieval
# --------------------------------------------------------------------------- #

def find_local_workbook(cfg: dict[str, Any]) -> Path | None:
    """Look for the synced copy of the workbook under the user's OneDrive folders.

    Saves having to hunt down the path by hand: once the library is synced, the
    file sits somewhere under a "OneDrive - <org>" folder in the home directory,
    but the exact location depends on which folder the shortcut was added for.
    """
    name = cfg.get("excel_filename") or "ADHOC.xlsx"
    hint = (cfg.get("excel_path_hint") or "").lower()

    try:
        roots = [p for p in Path.home().iterdir()
                 if p.is_dir() and any(k in p.name.lower()
                                       for k in ("onedrive", "sharepoint", "concentrix"))]
    except OSError:
        return None
    if not roots:
        return None

    matches: list[Path] = []
    deadline = time.monotonic() + 25
    for root in roots:
        try:
            for path in root.rglob(name):
                matches.append(path)
                if time.monotonic() > deadline:
                    break
        except OSError:
            continue
        if time.monotonic() > deadline:
            log("    (stopped searching after 25s)")
            break

    if not matches:
        return None
    preferred = [p for p in matches if hint and hint in str(p).lower()]
    return max(preferred or matches, key=lambda p: p.stat().st_mtime)


def local_workbook(cfg: dict[str, Any]) -> Path | None:
    """The workbook on disk, if there is one - configured path first, then a search."""
    configured = cfg.get("excel_path", "").strip()
    if configured:
        path = Path(configured).expanduser()
        if not path.exists():
            sys.exit(f"excel_path points at {path}, which does not exist")
        return path
    return find_local_workbook(cfg)


def resolve_workbook(cfg: dict[str, Any], page: Page | None) -> Path:
    found = local_workbook(cfg)
    if found:
        log(f"Reading {found}")
        return found

    if page is None:
        sys.exit(f"Could not find {cfg.get('excel_filename', 'the workbook')} on this PC, "
                 f"and no browser is available to download it")

    log("No local copy found - downloading the workbook from SharePoint")
    target = HERE / "ADHOC.xlsx"
    with page.expect_download(timeout=120_000) as dl:
        page.goto(cfg["excel_download_url"])
    dl.value.save_as(target)
    log(f"Downloaded to {target}")
    return target


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #

def launch_browser(pw, cfg: dict[str, Any]):
    """Reuse one browser profile so the SharePoint sign-in survives between runs."""
    common = dict(
        user_data_dir=str(PROFILE_DIR),
        headless=False,
        args=["--start-maximized"],
        no_viewport=True,
        accept_downloads=True,
    )
    channel = cfg.get("browser_channel", "msedge")
    if channel:
        try:
            return pw.chromium.launch_persistent_context(channel=channel, **common)
        except Exception as exc:
            log(f"Could not start '{channel}' ({exc}); falling back to bundled Chromium.")
    return pw.chromium.launch_persistent_context(**common)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Transfer audit rows into the Focus Audit list.")
    p.add_argument("--inspect", action="store_true",
                   help="capture the form's fields and dropdown options, then exit")
    p.add_argument("--dry-run", action="store_true",
                   help="fill the first row but stop before saving")
    p.add_argument("--preview", action="store_true",
                   help="print what would be submitted and what would be skipped, "
                        "without opening a browser")
    p.add_argument("--date", help="only rows whose auditDate matches (e.g. 08/01/2026)")
    p.add_argument("--last-days", type=int, default=7,
                   help="how far back to look when --date is not given (default 7)")
    p.add_argument("--all", action="store_true", help="ignore the date window entirely")
    p.add_argument("--limit", type=int, help="stop after this many rows")
    p.add_argument("--force", action="store_true",
                   help="include rows already recorded as submitted")
    p.add_argument("--yes", action="store_true", help="skip the confirmation prompt")
    p.add_argument("--include-unmatched", action="store_true",
                   help="attempt rows whose Selection Criteria has no matching option, "
                        "pausing so you can pick one, instead of skipping them")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    cfg = load_config()
    log(f"Focus Audit transfer v{VERSION}")

    if args.preview:
        found = local_workbook(cfg)
        if found:
            try:
                log(f"Reading {found}")
                rows = select_rows(cfg, read_rows(cfg, found), args)
                run_preview(cfg, *partition_rows(cfg, rows))
                return 0
            finally:
                flush_log()
        log("No local copy found - opening a browser to download the workbook")

    PROFILE_DIR.mkdir(exist_ok=True)

    with sync_playwright() as pw:
        context = launch_browser(pw, cfg)
        page = context.pages[0] if context.pages else context.new_page()

        try:
            if args.inspect:
                log("Inspecting the Focus Audit form...")
                inspect_form(page, cfg)
                return 0

            workbook = resolve_workbook(cfg, page)
            rows = select_rows(cfg, read_rows(cfg, workbook), args)

            if args.preview:
                run_preview(cfg, *partition_rows(cfg, rows))
                return 0

            skipped: list[tuple[Row, str, str]] = []
            if not args.include_unmatched:
                rows, skipped = partition_rows(cfg, rows)

            if skipped:
                log(f"\n{len(skipped)} row(s) will be SKIPPED - no matching option "
                    f"for a required dropdown:")
                for r, label, value in skipped:
                    log(f"  {r.audit_id}  {r.case_number}  {label} <- {value!r}")
                log("  (run with --include-unmatched to enter these by hand instead)")

            if not rows:
                log("\nNothing left to submit.")
                return 0

            log(f"\n{len(rows)} row(s) queued:")
            for r in rows:
                log(f"  {r.audit_id}  {r.audit_date}  {r.case_number}  "
                    f"{r.agent_email}  [{r.validation}]")

            if args.dry_run:
                rows = rows[:1]
                log("\nDry run - filling one entry and stopping before Save.")
            elif not args.yes:
                if input("\nProceed? [y/N] ").strip().lower() not in ("y", "yes"):
                    log("Cancelled.")
                    return 0

            ok = failed = 0
            for n, row in enumerate(rows, 1):
                log(f"\n[{n}/{len(rows)}] {row.audit_id} (Excel row {row.excel_row})")
                try:
                    form = open_new_item(page, cfg)
                    fill_row(form, row, cfg, interactive=True)
                    if args.dry_run:
                        log("  Filled. Check every field, then close the panel yourself.")
                        input("  Press Enter to finish the dry run... ")
                    else:
                        form.save()
                        mark_submitted(row.audit_id)
                        log("  Saved.")
                    ok += 1
                except (FormError, PWTimeout) as exc:
                    failed += 1
                    log(f"  FAILED: {exc}")
                    snap(page, f"row-{row.audit_id or row.excel_row}-failed")
                    if input("  Continue with the next row? [Y/n] ").strip().lower() in ("n", "no"):
                        break

            log(f"\nDone. {ok} succeeded, {failed} failed.")
            return 1 if failed else 0
        finally:
            flush_log()
            context.close()


if __name__ == "__main__":
    sys.exit(main())
