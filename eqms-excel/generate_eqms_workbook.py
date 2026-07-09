#!/usr/bin/env python3
"""Generate HP_Mainstream_EQMS.xlsx — a self-contained, formula-driven EQMS.

Design: the workbook lives in the team's SharePoint folder and is opened by all
auditors in Excel Online, where Microsoft 365 co-authoring provides the live
multi-user sync (no OneDrive client, no installs). Everything is plain Excel:
tables, INDEX/MATCH lookups, COUNTIFS dashboards, charts, data validation.

Sheets:
  START HERE            instructions
  Dashboard             executive KPIs + filters + charts
  Quick Case Dashboard  same, scoped to quick-case audits
  Audit Log             the database (Excel table, 2000 prepared rows)
  Roster                464 agents imported from the production roster
  Lists                 reason repositories, auditors, filter options, CC DL
  Calc / CalcQC         hidden aggregation engines feeding the dashboards
"""
import re
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROSTER_SRC = "/tmp/claude-0/-home-user-uigen/50079a4b-a861-5081-8afb-656977932e56/scratchpad/roster.xlsx"
OUT = "/home/user/uigen/eqms-excel/HP_Mainstream_EQMS.xlsx"

N_LOG = 2000          # prepared audit rows
YEAR = date.today().year

NAVY, RED, GREEN, BLUE, ORANGE = "1F3864", "C00000", "0A7D0A", "2A78D6", "C55A11"
LIGHT, MID = "F2F4F8", "D9E2F3"

INVALID_REASONS = [
    "Quick Case (Invalid)",
    "No Support Provided (Security Stop/Hold / Global Trade Fail)",
    "Cancel Per Customer / No Contact (CPC/NC)",
    "Case Voided (Test Case / Duplicate / Created by Mistake)",
    "Non-Tech Solution (General Inquiry / Warranty / How-To)",
    "Remote Solution (HW/SW Resolved Remotely via Troubleshooting)",
    "Offsite Solution (Bench / Depot Repair Service Order)",
    "Exchange Solution (Whole Unit Exchange Arranged)",
    "Onsite Solution (Break & Fix / Installation Work Order)",
    "Parts Shipped (CSR Part / RCD Media / Instant Ink PO)",
    "Referred - for Hardware Service (ASP / 3rd Party HW / DOA)",
    "Referred - for Software Support (Out of Scope / 3rd Party SW)",
    "Unresolved - Not Entitled for Support",
]
VALID_REASONS = [
    "Remote Solution (HW/SW Resolved Remotely via Troubleshooting)",
    "Parts Shipped (CSR Part / RCD Media / Instant Ink PO)",
    "Onsite Solution (Break & Fix / Installation Work Order)",
    "Offsite Solution (Bench / Depot Repair Service Order)",
    "Exchange Solution (Whole Unit Exchange Arranged)",
    "Non-Tech Solution (General Inquiry / Warranty / How-To)",
    "Cloud Recovery Download (Recovery Media with No Tech Concern)",
    "Referred - for Hardware Service (ASP / 3rd Party HW / DOA)",
    "Referred - for Software Support (Out of Scope / 3rd Party SW)",
    "Quick Case",
    "No Support Provided (Security Stop/Hold / Global Trade Fail)",
    "Cancel Per Customer / No Contact (CPC/NC)",
    "Case Voided (Test Case / Duplicate / Created by Mistake)",
    "Unresolved - Not Entitled for Support",
]
AUDITOR_EMAILS = [
    "sundeep.bhardwaj@concentrix.com", "ivy.serata@concentrix.com",
    "jonel.meniano1@concentrix.com", "archieval.decena@concentrix.com",
    "donnah.brual@concentrix.com", "emil.eduardo@concentrix.com",
    "maevenus.setias@concentrix.com", "kurtjohn.gatmaitan@concentrix.com",
    "attinareccie.padaya@concentrix.com", "christian.pila@concentrix.com",
    "shaileen.almodiel@concentrix.com", "darlene.presado@concentrix.com",
    "anita.fabales@concentrix.com", "jessalyn.sastrillo@concentrix.com",
    "caroline.chan@concentrix.com",
]
TIMEFRAMES = ["All Time", "Today", "Last 7 Days", "Last 14 Days", "Last 30 Days", "This Month"]

def email_to_name(e):
    local = e.split("@")[0]
    parts = [re.sub(r"\d+", "", p).capitalize() for p in local.split(".")]
    return " ".join(p for p in parts if p)

# ---------------------------------------------------------------- load roster
src = load_workbook(ROSTER_SRC, read_only=True)
rows = list(src.active.iter_rows(values_only=True))
hdr = [str(h or "").strip() for h in rows[0]]
agents = []
for r in rows[1:]:
    vals = ["" if v is None else str(v).strip() for v in r]
    if len(vals) < len(hdr):
        vals += [""] * (len(hdr) - len(vals))
    if vals[1]:
        agents.append(vals[:9])   # EmployeeNumber..SOMEmail
agents.sort(key=lambda a: a[1].lower())
tls = sorted({a[3] for a in agents if a[3]})
AUDITORS = [email_to_name(e) for e in AUDITOR_EMAILS]
print(f"roster: {len(agents)} agents, {len(tls)} team leaders")

# ---------------------------------------------------------------- styles
f_title = Font(bold=True, size=16, color="FFFFFF")
f_h = Font(bold=True, size=11, color="FFFFFF")
f_lbl = Font(bold=True, size=9, color="5A6B85")
f_kpi = Font(bold=True, size=22, color=NAVY)
f_kpi_md = Font(bold=True, size=12, color=NAVY)
fill_navy = PatternFill("solid", start_color=NAVY)
fill_red = PatternFill("solid", start_color=RED)
fill_light = PatternFill("solid", start_color=LIGHT)
fill_mid = PatternFill("solid", start_color=MID)
thin = Side(style="thin", color="B4C0D6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

def style_block(ws, rng, fill=None, font=None, align=None, bdr=True):
    for row in ws[rng]:
        for c in row:
            if fill: c.fill = fill
            if font: c.font = font
            if align: c.alignment = align
            if bdr: c.border = border

wb = Workbook()

# ---------------------------------------------------------------- START HERE
ws = wb.active
ws.title = "START HERE"
ws.sheet_properties.tabColor = NAVY
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 110
ws.merge_cells("B2:B3")
ws["B2"] = "🛡  HP MAINSTREAM EQMS — Excel Edition (live team co-authoring)"
ws["B2"].font = Font(bold=True, size=18, color="FFFFFF")
ws["B2"].fill = fill_navy
ws["B2"].alignment = center
lines = [
    "",
    "HOW THE LIVE SYNC WORKS",
    "This workbook lives in the team's SharePoint 'EQMS' folder. Everyone opens it IN THE BROWSER",
    "(SharePoint > click the file > it opens in Excel Online). Microsoft 365 co-authoring shows every",
    "auditor's entries to everyone else within seconds, automatically. Do NOT download personal copies.",
    "",
    "HOW TO LOG AN AUDIT  (sheet: Audit Log)",
    "  1. Go to the first empty row and enter the Audit Date.",
    "  2. Pick the Agent Name from the dropdown — EID, emails, TL, OM, SOM auto-fill instantly.",
    "  3. Fill Channel (if blank), Case Number, Genesys Transaction ID.",
    "  4. Pick 'Validation Result / Case Resolution': 'Valid (Pass)' or the invalid defect reason.",
    "  5. Pick the Expected Resolution Code, type your Remarks, pick your Auditor Name.",
    "  6. If the audit is INVALID, the '✉ Send Email' cell turns into a link — click it to open a",
    "     pre-addressed Outlook draft to the TL + OM (Cc: QA DL). Review and press Send.",
    "",
    "DASHBOARDS",
    "  'Dashboard' and 'Quick Case Dashboard' update live. Use the yellow filter cells at the top",
    "  (Timeframe / Team Leader / Auditor / Validation). Charts and KPIs follow the filters.",
    "",
    "RULES (please respect them — the formulas depend on it)",
    "  •  Only type in the WHITE columns of Audit Log. Grey columns are automatic — don't overwrite.",
    "  •  Never delete or insert rows in Audit Log; to void an audit, clear its white cells.",
    "  •  Audit Date is required — audits without a date don't count on the dashboards.",
    "  •  Disputes: the auditor (or Sundeep) edits the audit's row directly; co-authoring shows who's editing.",
    "  •  Admin (Sundeep) maintains the 'Roster' and 'Lists' sheets. 2,000 audit rows are prepared;",
    "     ask the admin to extend the table when you approach the end.",
    "  •  Channel and Internal LOB come from the Roster: the admin sets each agent's CHANNEL",
    "     (dropdown on the Roster sheet) and Internal LOB — audits then auto-fill them.",
    "",
    "VERSION HISTORY / BACKUP: SharePoint keeps automatic version history (file > Version History),",
    "so any mistake can be rolled back. No manual backups needed.",
]
r = 4
for t in lines:
    ws.cell(row=r, column=2, value=t)
    ws.cell(row=r, column=2).font = Font(bold=t.isupper() or t.startswith(("HOW", "DASH", "RULES", "VERSION")), size=11,
                                         color=NAVY if (t and (t[0].isalpha() and t.isupper())) else "222222")
    ws.cell(row=r, column=2).alignment = left
    r += 1

# ---------------------------------------------------------------- Lists sheet
lists = wb.create_sheet("Lists")
lists.sheet_properties.tabColor = "888888"
lists["A1"], lists["B1"], lists["C1"], lists["D1"], lists["E1"], lists["F1"], lists["G1"], lists["H1"] = (
    "Case Resolution options", "Expected Resolution Codes", "Auditors", "Timeframes",
    "CC distribution list (edit)", "Validation filter", "Team Leaders", "Agent Names")
for c in "ABCDEFGH":
    lists[f"{c}1"].font = Font(bold=True)
    lists.column_dimensions[c].width = 46
res_options = ["Valid (Pass)"] + INVALID_REASONS
for i, v in enumerate(res_options, start=2): lists.cell(row=i, column=1, value=v)
for i, v in enumerate(VALID_REASONS, start=2): lists.cell(row=i, column=2, value=v)
for i, v in enumerate(AUDITORS, start=2): lists.cell(row=i, column=3, value=v)
for i, v in enumerate(TIMEFRAMES, start=2): lists.cell(row=i, column=4, value=v)
lists["E2"] = "PH_QUE_NE_HP_Quality_Support@concentrix.com"
for i, v in enumerate(["All Statuses", "Valid", "Invalid"], start=2): lists.cell(row=i, column=6, value=v)
lists["G2"] = "All Team Leaders"
for i, v in enumerate(tls, start=3): lists.cell(row=i, column=7, value=v)
lists["I1"] = "Auditor filter"; lists["I1"].font = Font(bold=True)
lists["I2"] = "All Auditors"
for i, v in enumerate(AUDITORS, start=3): lists.cell(row=i, column=9, value=v)

# ---------------------------------------------------------------- Roster
ros = wb.create_sheet("Roster")
ros.sheet_properties.tabColor = GREEN
ROS_HDR = ["EmployeeNumber", "FullName", "Email", "ImmediateSupervisor", "ImmediateSupervisorEmail",
           "Manager", "ManagerEmail", "SOM", "SOMEmail", "CHANNEL", "Internal LOB"]
ros.append(ROS_HDR)
for a in agents:
    # CHANNEL is admin-maintained (dropdown below); Internal LOB defaults org-wide
    ros.append(a + ["", "HP Mainstream"])
n_ros = len(agents) + 1
for i, w in enumerate([14, 34, 36, 30, 36, 30, 36, 30, 36, 14, 16], start=1):
    ros.column_dimensions[get_column_letter(i)].width = w
t = Table(displayName="tblRoster", ref=f"A1:K{n_ros}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ros.add_table(t)
ros.freeze_panes = "A2"
# CHANNEL dropdown (suggestions, not enforced) for the admin to fill per agent
lists["K1"] = "Channels"; lists["K1"].font = Font(bold=True)
for i, v in enumerate(["Voice", "Chat", "Email", "Back Office"], start=2):
    lists.cell(row=i, column=11, value=v)
dv_ch = DataValidation(type="list", formula1="=Lists!$K$2:$K$5", allow_blank=True, showErrorMessage=False)
ros.add_data_validation(dv_ch)
dv_ch.add(f"J2:J{n_ros}")

# agent names into Lists!H for the dropdown named range
for i, a in enumerate(agents, start=2):
    lists.cell(row=i, column=8, value=a[1])

# ---------------------------------------------------------------- Audit Log
log = wb.create_sheet("Audit Log")
log.sheet_properties.tabColor = BLUE
HDR = ["auditId", "auditDate", "agentName", "agentEid", "agentEmail", "teamLeader", "tlEmail",
       "operationsManager", "omEmail", "somEmail", "channel", "lob", "caseNumber",
       "genesysTransactionId", "validation", "caseResolution", "expectedResolutionCode",
       "remarks", "auditorName", "quickCase", "sendEmail"]
log.append(HDR)
LAST = N_LOG + 1
# Plain relative references (not [@col] structured refs): identical behaviour in
# Excel, and verifiable in LibreOffice which mis-parses the [@col] shorthand.
def LK(col, r):
    # empty roster cells must come back as "" (bare INDEX turns them into 0)
    idx = f'INDEX(Roster!${col}:${col},MATCH($C{r},Roster!$B:$B,0))'
    return f'IFERROR(IF({idx}="","",{idx}&""),"—")'
def mail_body(r):
    return (
        f'"Hi "&$F{r}&","&CHAR(10)&"Good day!"&CHAR(10)&CHAR(10)'
        f'&"A quality monitoring audit has been completed for "&$C{r}&"."&CHAR(10)&CHAR(10)'
        f'&"QUALITY AUDIT DETAILS"&CHAR(10)'
        f'&"Agent Name: "&$C{r}&" (EID "&$D{r}&")"&CHAR(10)'
        f'&"Audit Date: "&TEXT($B{r},"yyyy-mm-dd")&CHAR(10)'
        f'&"Team Leader: "&$F{r}&CHAR(10)'
        f'&"Case Number: "&$M{r}&CHAR(10)'
        f'&"Transaction ID: "&$N{r}&CHAR(10)'
        f'&"Channel / LOB: "&$K{r}&" / "&$L{r}&CHAR(10)'
        f'&"Evaluator: "&$S{r}&CHAR(10)&CHAR(10)'
        f'&"INVALID QUALITY AUDIT DETAILS"&CHAR(10)'
        f'&"Validation: Invalid"&CHAR(10)'
        f'&"Invalid Reason: "&$P{r}&CHAR(10)'
        f'&"Expected Resolution Code: "&$Q{r}&CHAR(10)'
        f'&"Remarks: "&LEFT($R{r},600)&CHAR(10)&CHAR(10)'
        f'&"This is an automated notification from HP Mainstream EQMS."'
    )
def formulas_for(r):
    return {
        "auditId": f'=IF($C{r}="","","AUD-{YEAR}-"&TEXT(ROW()-1,"000000"))',
        "agentEid": f'=IF($C{r}="","",{LK("A", r)})',
        "agentEmail": f'=IF($C{r}="","",{LK("C", r)})',
        "teamLeader": f'=IF($C{r}="","",{LK("D", r)})',
        "tlEmail": f'=IF($C{r}="","",{LK("E", r)})',
        "operationsManager": f'=IF($C{r}="","",{LK("F", r)})',
        "omEmail": f'=IF($C{r}="","",{LK("G", r)})',
        "somEmail": f'=IF($C{r}="","",{LK("I", r)})',
        "channel": f'=IF($C{r}="","",{LK("J", r)})',
        "lob": f'=IF($C{r}="","",{LK("K", r)})',
        "validation": f'=IF($C{r}="","",IF($P{r}="","—",'
                      f'IF($P{r}="Valid (Pass)","Valid","Invalid")))',
        "quickCase": f'=IF($C{r}="","",IF(OR(ISNUMBER(SEARCH("quick case",$P{r})),'
                     f'ISNUMBER(SEARCH("quick case",$Q{r}))),"Yes","No"))',
        "sendEmail": f'=IF($O{r}="Invalid",HYPERLINK("mailto:"&$G{r}&";"&$I{r}'
                     f'&"?cc="&Lists!$E$2'
                     f'&"&subject="&_xlfn.ENCODEURL("[INVALID] Quality audit monitoring completed for "&UPPER($C{r}))'
                     f'&"&body="&_xlfn.ENCODEURL({mail_body(r)}),"✉ Open Email Draft"),"")',
    }
col_idx = {h: i + 1 for i, h in enumerate(HDR)}
for r in range(2, LAST + 1):
    for h, f in formulas_for(r).items():
        log.cell(row=r, column=col_idx[h], value=f)
t = Table(displayName="tblAudits", ref=f"A1:U{LAST}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium16", showRowStripes=True)
log.add_table(t)
log.freeze_panes = "D2"
widths = [16, 12, 32, 12, 32, 28, 32, 28, 32, 32, 12, 14, 14, 22, 10, 40, 40, 50, 22, 10, 18]
for i, w in enumerate(widths, start=1):
    log.column_dimensions[get_column_letter(i)].width = w
# grey the automatic columns, keep entry columns white
auto_cols = ["A", "D", "E", "F", "G", "H", "I", "J", "O", "T", "U"]
autofill = PatternFill("solid", start_color="EDEFF3")
for c in auto_cols:
    for r in range(2, LAST + 1):
        log[f"{c}{r}"].fill = autofill
for r in range(2, LAST + 1):
    log.cell(row=r, column=2).number_format = "yyyy-mm-dd"
# conditional formatting
red_font = Font(color=RED, bold=True); green_font = Font(color=GREEN, bold=True)
log.conditional_formatting.add(f"O2:O{LAST}",
    CellIsRule(operator="equal", formula=['"Invalid"'], font=red_font,
               fill=PatternFill("solid", start_color="FBE4E4")))
log.conditional_formatting.add(f"O2:O{LAST}",
    CellIsRule(operator="equal", formula=['"Valid"'], font=green_font,
               fill=PatternFill("solid", start_color="E4F5E4")))
log.conditional_formatting.add(f"M2:M{LAST}",
    FormulaRule(formula=[f'AND($M2<>"",COUNTIF($M$2:$M${LAST},$M2)>1)'],
                fill=PatternFill("solid", start_color="FFE49C")))
# data validations
def dv_list(name, target):
    d = DataValidation(type="list", formula1=f"={name}", allow_blank=True, showErrorMessage=True)
    log.add_data_validation(d); d.add(target)
dv_list("AgentNames", f"C2:C{LAST}")
dv_list("ResOptions", f"P2:P{LAST}")
dv_list("ExpOptions", f"Q2:Q{LAST}")
dv_list("AuditorNames", f"S2:S{LAST}")

# ---------------------------------------------------------------- named ranges
def add_name(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
add_name("AgentNames", f"Lists!$H$2:$H${len(agents)+1}")
add_name("ResOptions", f"Lists!$A$2:$A${len(res_options)+1}")
add_name("ExpOptions", f"Lists!$B$2:$B${len(VALID_REASONS)+1}")
add_name("AuditorNames", f"Lists!$C$2:$C${len(AUDITORS)+1}")
add_name("Timeframes", f"Lists!$D$2:$D${len(TIMEFRAMES)+1}")
add_name("ValFilter", "Lists!$F$2:$F$4")
add_name("TLFilter", f"Lists!$G$2:$G${len(tls)+2}")
add_name("AudFilter", f"Lists!$I$2:$I${len(AUDITORS)+2}")

# ---------------------------------------------------------------- calc engine
def build_calc(name, dash, qc_only):
    """Hidden sheet computing everything a dashboard needs from its filters."""
    cs = wb.create_sheet(name)
    cs.sheet_state = "hidden"
    D = f"'{dash}'"
    LOGC = lambda c: f"'Audit Log'!${c}$2:${c}${LAST}"
    qc_pair = f",{LOGC('T')},\"Yes\"" if qc_only else ""
    base = (f"{LOGC('B')},\">=\"&$B$2,{LOGC('F')},$B$3,{LOGC('S')},$B$4" + qc_pair)
    cs["A1"] = "filter engine"
    cs["A2"], cs["A3"], cs["A4"], cs["A5"] = "startDate", "tlCrit", "auditorCrit", "valCrit"
    cs["B2"] = (f'=IF({D}!$C$2="All Time",0,IF({D}!$C$2="Today",TODAY(),'
                f'IF({D}!$C$2="Last 7 Days",TODAY()-6,IF({D}!$C$2="Last 14 Days",TODAY()-13,'
                f'IF({D}!$C$2="Last 30 Days",TODAY()-29,DATE(YEAR(TODAY()),MONTH(TODAY()),1))))))')
    cs["B3"] = f'=IF({D}!$E$2="All Team Leaders","*",{D}!$E$2)'
    cs["B4"] = f'=IF({D}!$G$2="All Auditors","*",{D}!$G$2)'
    cs["B5"] = f'=IF({D}!$I$2="All Statuses","*",{D}!$I$2)'
    # headline counts
    cs["B7"] = f"=COUNTIFS({base},{LOGC('O')},$B$5)"                    # filtered
    cs["B8"] = f"=COUNTIFS({base},{LOGC('O')},\"Valid\",{LOGC('O')},$B$5)"
    cs["B9"] = f"=COUNTIFS({base},{LOGC('O')},\"Invalid\",{LOGC('O')},$B$5)"
    cs["B10"] = "=IFERROR($B$9/$B$7,0)"
    cs["B11"] = (f"=IFERROR(TEXT(_xlfn.MAXIFS({LOGC('B')},{LOGC('B')},\">=\"&$B$2,"
                 f"{LOGC('F')},$B$3,{LOGC('S')},$B$4{qc_pair},{LOGC('O')},$B$5),\"mmm d, yyyy\"),\"—\")")
    # 14-day trend  D2:E15
    cs["D1"], cs["E1"] = "day", "audits"
    for i in range(14):
        r = i + 2
        cs.cell(row=r, column=4, value=f"=TODAY()-{13-i}").number_format = "mm-dd"
        cs.cell(row=r, column=5,
                value=f"=COUNTIFS({LOGC('B')},$D{r},{LOGC('F')},$B$3,{LOGC('S')},$B$4"
                      f"{qc_pair},{LOGC('O')},$B$5)")
    # defective code distribution  G:I, top6 K:M
    cs["G1"], cs["H1"], cs["I1"] = "invalid reason", "defects", "rank helper"
    for i, reason in enumerate(INVALID_REASONS):
        r = i + 2
        cs.cell(row=r, column=7, value=reason)
        cs.cell(row=r, column=8, value=f"=COUNTIFS({base},{LOGC('P')},$G{r})")
        cs.cell(row=r, column=9, value=f"=$H{r}+ROW()/100000")
    nR = len(INVALID_REASONS) + 1
    cs["K1"], cs["L1"] = "top code", "defects"
    for k in range(6):
        r = k + 2
        cs.cell(row=r, column=11,
                value=f"=IFERROR(INDEX($G$2:$G${nR},MATCH(LARGE($I$2:$I${nR},{k+1}),$I$2:$I${nR},0)),\"\")")
        cs.cell(row=r, column=12, value=f"=IFERROR(INT(LARGE($I$2:$I${nR},{k+1})),0)")
    # TL performance  N:R, top5 T:V
    cs["N1"], cs["O1"], cs["P1"], cs["Q1"], cs["R1"] = "TL", "audits", "valid", "rank", "pass"
    for i, tl in enumerate(tls):
        r = i + 2
        cs.cell(row=r, column=14, value=tl)
        cs.cell(row=r, column=15,
                value=f"=COUNTIFS({LOGC('B')},\">=\"&$B$2,{LOGC('F')},$N{r},{LOGC('S')},$B$4{qc_pair})")
        cs.cell(row=r, column=16,
                value=f"=COUNTIFS({LOGC('B')},\">=\"&$B$2,{LOGC('F')},$N{r},{LOGC('S')},$B$4"
                      f"{qc_pair},{LOGC('O')},\"Valid\")")
        cs.cell(row=r, column=17, value=f"=$O{r}+ROW()/100000")
        cs.cell(row=r, column=18, value=f"=IFERROR($P{r}/$O{r},0)")
    nT = len(tls) + 1
    cs["T1"], cs["U1"], cs["V1"] = "top TL", "audits", "pass %"
    for k in range(5):
        r = k + 2
        cs.cell(row=r, column=20,
                value=f"=IFERROR(INDEX($N$2:$N${nT},MATCH(LARGE($Q$2:$Q${nT},{k+1}),$Q$2:$Q${nT},0)),\"\")")
        cs.cell(row=r, column=21, value=f"=IFERROR(INT(LARGE($Q$2:$Q${nT},{k+1})),0)")
        cs.cell(row=r, column=22,
                value=f"=IFERROR(INDEX($R$2:$R${nT},MATCH(LARGE($Q$2:$Q${nT},{k+1}),$Q$2:$Q${nT},0)),0)")
        cs.cell(row=r, column=22).number_format = "0.0%"
    # top agents with defects (QC dashboard uses this)  X:Z top6 AB:AC
    cs["X1"], cs["Y1"], cs["Z1"] = "agent", "defects", "rank"
    for i, a in enumerate(agents):
        r = i + 2
        cs.cell(row=r, column=24, value=a[1])
        cs.cell(row=r, column=25,
                value=f"=COUNTIFS({base},{LOGC('C')},$X{r},{LOGC('O')},\"Invalid\")")
        cs.cell(row=r, column=26, value=f"=$Y{r}+ROW()/100000")
    nA = len(agents) + 1
    cs["AB1"], cs["AC1"] = "top agent", "defects"
    for k in range(6):
        r = k + 2
        cs.cell(row=r, column=28,
                value=f"=IFERROR(INDEX($X$2:$X${nA},MATCH(LARGE($Z$2:$Z${nA},{k+1}),$Z$2:$Z${nA},0)),\"\")")
        cs.cell(row=r, column=29, value=f"=IFERROR(INT(LARGE($Z$2:$Z${nA},{k+1})),0)")
    return cs

# ---------------------------------------------------------------- dashboards
def build_dashboard(title, calc, qc):
    d = wb.create_sheet(title)
    d.sheet_properties.tabColor = ORANGE if qc else BLUE
    d.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJK", [2, 14, 22, 13, 26, 10, 24, 12, 14, 16, 16]):
        d.column_dimensions[col].width = w
    d.merge_cells("B1:K1")
    d["B1"] = ("⚡ QUICK CASE DASHBOARD — live" if qc else "🛡 EXECUTIVE DASHBOARD — live") + \
              "  (co-authoring: everyone sees the same numbers)"
    d["B1"].font = f_title; d["B1"].fill = fill_navy; d["B1"].alignment = center
    d.row_dimensions[1].height = 30
    # filter row (yellow editable cells)
    yellow = PatternFill("solid", start_color="FFF2CC")
    filters = [("B2", "📅 Timeframe", "C2", "Timeframes", "All Time"),
               ("D2", "👤 Team Leader", "E2", "TLFilter", "All Team Leaders"),
               ("F2", "🏷 Auditor", "G2", "AudFilter", "All Auditors"),
               ("H2", "✅ Validation", "I2", "ValFilter", "All Statuses")]
    for lbl_cell, lbl, val_cell, list_name, default in filters:
        d[lbl_cell] = lbl; d[lbl_cell].font = Font(bold=True, size=10)
        d[val_cell] = default; d[val_cell].fill = yellow; d[val_cell].border = border
        dv = DataValidation(type="list", formula1=f"={list_name}", allow_blank=False)
        d.add_data_validation(dv); dv.add(val_cell)
    d.row_dimensions[2].height = 22
    # KPI cards
    C = f"'{calc.title}'"
    kpis = [("B4", "C5", "FILTERED AUDITS" if not qc else "QUICK CASE AUDITS", f"={C}!$B$7", None, NAVY),
            ("D4", "E5", "VALID AUDITS", f"={C}!$B$8", None, GREEN),
            ("F4", "G5", "INVALID AUDITS", f"={C}!$B$9", None, RED),
            ("H4", "I5", "INVALID PERCENTAGE", f"={C}!$B$10", "0.0%", RED)]
    for lbl_cell, val_cell, lbl, f, fmt, color in kpis:
        d[lbl_cell] = lbl; d[lbl_cell].font = f_lbl
        d[val_cell] = f; d[val_cell].font = Font(bold=True, size=20, color=color)
        if fmt: d[val_cell].number_format = fmt
    d["B7"] = "TOP DEFECTIVE RESOLUTION CODE"; d["B7"].font = f_lbl
    d.merge_cells("B8:E8")
    d["B8"] = f'=IF({C}!$L$2=0,"—",{C}!$K$2&"  ("&{C}!$L$2&"x)")'
    d["B8"].font = Font(bold=True, size=12, color=RED)
    d["F7"] = "LATEST AUDIT (FILTERED)"; d["F7"].font = f_lbl
    d.merge_cells("F8:H8")
    d["F8"] = f"={C}!$B$11"; d["F8"].font = f_kpi_md
    # charts
    cn = calc.title
    trend = LineChart(); trend.title = "Audit Trend (Last 14 Days)"; trend.height = 7.2; trend.width = 15
    data = Reference(calc, min_col=5, min_row=1, max_row=15)
    cats = Reference(calc, min_col=4, min_row=2, max_row=15)
    trend.add_data(data, titles_from_data=True); trend.set_categories(cats)
    trend.series[0].graphicalProperties.line.solidFill = BLUE
    trend.series[0].graphicalProperties.line.width = 22000
    trend.legend = None
    d.add_chart(trend, "B10")
    dough = DoughnutChart(); dough.title = "Valid vs Invalid"; dough.height = 7.2; dough.width = 8.5
    dd = Reference(calc, min_col=2, min_row=8, max_row=9)
    dl_start = 8
    dough.add_data(dd)
    labels_ref = Reference(calc, min_col=1, min_row=8, max_row=9)
    calc["A8"], calc["A9"] = "Valid", "Invalid"
    dough.set_categories(labels_ref)
    s = dough.series[0]
    for idx, colr in ((0, GREEN), (1, RED)):
        pt = DataPoint(idx=idx); pt.graphicalProperties.solidFill = colr
        s.data_points.append(pt)
    d.add_chart(dough, "H10")
    bar1 = BarChart(); bar1.type = "bar"; bar1.title = "Top Defective Case Resolution Codes"
    bar1.height = 7.5; bar1.width = 15
    b1d = Reference(calc, min_col=12, min_row=1, max_row=7)
    b1c = Reference(calc, min_col=11, min_row=2, max_row=7)
    bar1.add_data(b1d, titles_from_data=True); bar1.set_categories(b1c)
    bar1.series[0].graphicalProperties.solidFill = RED
    bar1.legend = None
    d.add_chart(bar1, "B26")
    if qc:
        bar2 = BarChart(); bar2.type = "bar"; bar2.title = "Agents with Most Quick Case Defects"
        b2d = Reference(calc, min_col=29, min_row=1, max_row=7)
        b2c = Reference(calc, min_col=28, min_row=2, max_row=7)
    else:
        bar2 = BarChart(); bar2.type = "bar"; bar2.title = "Audits by Team Leader (Top 5)"
        b2d = Reference(calc, min_col=21, min_row=1, max_row=6)
        b2c = Reference(calc, min_col=20, min_row=2, max_row=6)
    bar2.height = 7.5; bar2.width = 15
    bar2.add_data(b2d, titles_from_data=True); bar2.set_categories(b2c)
    bar2.series[0].graphicalProperties.solidFill = BLUE if not qc else RED
    bar2.legend = None
    d.add_chart(bar2, "H26")
    d.sheet_view.zoomScale = 90
    return d

calc_main = build_calc("Calc", "Dashboard", qc_only=False)
calc_qc = build_calc("CalcQC", "Quick Case Dashboard", qc_only=True)
build_dashboard("Dashboard", calc_main, qc=False)
build_dashboard("Quick Case Dashboard", calc_qc, qc=True)

# sheet order: START HERE, Dashboard, QC Dashboard, Audit Log, Roster, Lists, (hidden calcs)
order = ["START HERE", "Dashboard", "Quick Case Dashboard", "Audit Log", "Roster", "Lists", "Calc", "CalcQC"]
wb._sheets = [wb[n] for n in order]

wb.save(OUT)
print("saved", OUT)
